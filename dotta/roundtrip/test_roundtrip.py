import io
import json
import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

import betteroffice_pptx as bo
import make_deck

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FONT = ROOT / "crates" / "ooxml-text" / "tests" / "fonts" / "LiberationSans-Regular.ttf"
NOTO_JP = ROOT / "packages" / "fonts-cjk" / "assets" / "NotoSansJP-Regular.otf"
WASM = ROOT / "packages" / "pptx" / "src" / "wasm" / "generated" / "pptx_wasm_bg.wasm"
SHARED_PARTS = [
    "ppt/presentation.xml",
    "ppt/slideMasters/slideMaster1.xml",
    "ppt/theme/theme1.xml",
    *[f"ppt/slideLayouts/slideLayout{index}.xml" for index in range(1, 12)],
]


@pytest.fixture(scope="module")
def source(tmp_path_factory) -> bytes:
    path = tmp_path_factory.mktemp("deck") / "source.pptx"
    make_deck.build(str(path))
    return path.read_bytes()


def parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def fonts(data: bytes) -> list[tuple[str, str | None, str | None]]:
    prs = Presentation(io.BytesIO(data))
    found = []
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        rpr = run._r.find(qn("a:rPr"))
                        latin = None if rpr is None else rpr.find(qn("a:latin"))
                        ea = None if rpr is None else rpr.find(qn("a:ea"))
                        found.append((
                            run.text,
                            None if latin is None else latin.get("typeface"),
                            None if ea is None else ea.get("typeface"),
                        ))
    return found


def table_shape(data: bytes):
    slide = Presentation(io.BytesIO(data)).slides[make_deck.TABLE_SLIDE]
    return next(shape for shape in slide.shapes if shape.has_table)


def merges(data: bytes):
    table = table_shape(data).table
    return [
        [(cell.is_merge_origin, cell.is_spanned, cell.span_width, cell.span_height) for cell in row.cells]
        for row in table.rows
    ]


def table_sizes(data: bytes):
    shape = table_shape(data)
    return (
        [column.width for column in shape.table.columns],
        [row.height for row in shape.table.rows],
        shape.width,
        shape.height,
    )


def open_deck(data: bytes) -> bo.Presentation:
    deck = bo.Presentation.open(data)
    deck.register_font("Liberation Sans", FONT.read_bytes())
    deck.register_fallback_font("Noto Sans JP", NOTO_JP.read_bytes())
    return deck


def geometries(display_list) -> list[str]:
    return [primitive["geometry"] for primitive in display_list.to_dict()["primitives"] if primitive["kind"] == "shape"]


@pytest.fixture(scope="module")
def edited_by_python(source) -> bytes:
    deck = open_deck(source)
    title = next(shape for shape in deck.slide(0).shapes if shape.placeholder in ("ctrTitle", "title"))
    deck.insert_text(title.stories[0].id, 0, "改訂: ", font_family="游ゴシック", bold=True)
    return deck.save()


def test_only_the_edited_slide_changes(source, edited_by_python):
    before, after = parts(source), parts(edited_by_python)
    assert before.keys() == after.keys()
    changed = sorted(name for name in before if before[name] != after[name])
    assert changed == ["ppt/slides/slide1.xml"]


def test_font_names_survive_and_new_text_keeps_the_run_face(source, edited_by_python):
    before, after = fonts(source), fonts(edited_by_python)
    assert after[0] == ("改訂: 2026年度 第3四半期 売上報告", "游ゴシック", "游ゴシック")
    assert after[1:] == before[1:]
    assert b"Noto Sans JP" not in b"".join(parts(edited_by_python).values())


def test_the_title_still_inherits_its_frame_from_the_layout(edited_by_python):
    slide = Presentation(io.BytesIO(edited_by_python)).slides[0]
    title = slide.shapes.title._element
    assert title.spPr.find(qn("a:xfrm")) is None
    assert title.spPr.find(qn("a:prstGeom")) is None


def test_merged_cells_and_table_sizes_survive(source, edited_by_python):
    assert merges(edited_by_python) == merges(source)
    assert table_sizes(edited_by_python) == table_sizes(source)
    assert table_sizes(source)[:2] == (make_deck.COLUMN_WIDTHS, make_deck.ROW_HEIGHTS)


def test_render_png_draws_a_layout_ellipse_placeholder_as_an_ellipse(source):
    deck = open_deck(source)
    body = next(shape for shape in deck.slide(make_deck.ELLIPSE_SLIDE).shapes if shape.placeholder != "title")
    assert body.geometry == "rect"
    assert "ellipse" in geometries(deck.render_slide(make_deck.ELLIPSE_SLIDE))
    assert deck.render_png(make_deck.ELLIPSE_SLIDE).bytes[:8] == b"\x89PNG\r\n\x1a\n"


def run_wasm_edit(source: bytes, tmp_path: Path, widths, heights) -> tuple[bytes, dict]:
    bun = shutil.which("bun")
    if not bun or not WASM.exists():
        if os.environ.get("DOTTA_REQUIRE_WASM"):
            pytest.fail("bun and the built pptx wasm are required")
        pytest.skip("bun or the built pptx wasm is missing")
    source_path, output_path = tmp_path / "in.pptx", tmp_path / "out.pptx"
    source_path.write_bytes(source)
    sizes = json.dumps({"slide": make_deck.TABLE_SLIDE, "widths": widths, "heights": heights})
    result = subprocess.run(
        [bun, str(HERE / "wasm_edit.ts"), str(source_path), str(output_path), sizes, str(make_deck.ELLIPSE_SLIDE)],
        cwd=ROOT, check=True, capture_output=True, text=True,
    )
    return output_path.read_bytes(), json.loads(result.stdout)


def test_the_editor_engine_resizes_a_table_and_draws_the_ellipse(source, tmp_path):
    widths = [3000000, 2000000, 2500000, 3000000]
    heights = [600000, 500000, 800000, 500000, 900000]
    edited, report = run_wasm_edit(source, tmp_path, widths, heights)
    assert "ellipse" in report["geometries"]
    assert table_sizes(edited) == (widths, heights, sum(widths), sum(heights))
    assert merges(edited) == merges(source)
    assert fonts(edited) == fonts(source)
    before, after = parts(source), parts(edited)
    assert sorted(name for name in before if before[name] != after.get(name)) == ["ppt/slides/slide2.xml"]
    for name in SHARED_PARTS:
        assert after[name] == before[name], name
